from pathlib import Path
import csv
import json
import os
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
FILES = ROOT / "files"

def text_examples():
    path = FILES / "reading_file_example.txt"

    with open(path, encoding="utf-8") as f:
        print("read():")
        print(f.read())

    with open(path, encoding="utf-8") as f:
        print("read(10):", repr(f.read(10)))

    with open(path, encoding="utf-8") as f:
        print("readline():", repr(f.readline()))

    with open(path, encoding="utf-8") as f:
        print("readlines():", f.readlines())

    with open(path, encoding="utf-8") as f:
        print("splitlines():", f.read().splitlines())

def write_examples():
    path = FILES / "writing_file_example.txt"

    with open(path, "w", encoding="utf-8") as f:
        f.write("This file was created using write mode.\n")

    with open(path, "a", encoding="utf-8") as f:
        f.write("This line was appended using append mode.\n")

    print("Written:", path)

def json_examples():
    person = {
        "name": "Asabeneh",
        "country": "Finland",
        "city": "Helsinki",
        "skills": ["JavaScript", "React", "Python"]
    }

    text = json.dumps(person, indent=4)
    print("\nJSON string:")
    print(text)

    restored = json.loads(text)
    print("Restored dictionary:", restored)

    output = FILES / "json_example.json"
    with open(output, "w", encoding="utf-8") as f:
        json.dump(person, f, ensure_ascii=False, indent=4)

def csv_examples():
    path = FILES / "csv_example.csv"
    with open(path, encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        print("\nCSV rows:")
        for row in reader:
            print(dict(row))

def xlsx_examples():
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        print("\nXLSX: install dependencies with `pip install -r requirements.txt`")
        return

    path = FILES / "sample.xlsx"

    if not path.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Students"
        ws.append(["Name", "Course", "Score"])
        ws.append(["Aarav", "MCA", 88])
        ws.append(["Meera", "MCA", 92])
        wb.save(path)

    wb = load_workbook(path, read_only=True)
    ws = wb.active
    print("\nXLSX rows:")
    for row in ws.iter_rows(values_only=True):
        print(row)
    wb.close()

def xml_examples():
    path = FILES / "xml_example.xml"
    tree = ET.parse(path)
    root = tree.getroot()

    print("\nXML:")
    print("Root tag:", root.tag)
    print("Attributes:", root.attrib)

    for child in root:
        print("Child:", child.tag)

def deletion_example():
    path = FILES / "_temporary_file.txt"
    path.write_text("temporary", encoding="utf-8")

    if os.path.exists(path):
        os.remove(path)
        print("\nDeleted:", path.name)

if __name__ == "__main__":
    print("=== TXT ===")
    text_examples()
    print("\n=== WRITE / APPEND ===")
    write_examples()
    print("\n=== JSON ===")
    json_examples()
    print("\n=== CSV ===")
    csv_examples()
    print("\n=== XLSX ===")
    xlsx_examples()
    print("\n=== XML ===")
    xml_examples()
    print("\n=== DELETE ===")
    deletion_example()
